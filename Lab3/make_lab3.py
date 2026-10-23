#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Генератор отчёта по лабораторной работе №3 «Доработка требований» (Вариант 3, «File Searcher»).
Дисциплина: Разработка и анализ требований к программному обеспечению (БГУИР)
Студент: Забелич П. А., гр. 451001. Преподаватель: Данилова Г. В.

Оформление (титульный лист, шрифты, нумерованные списки) берётся из docx-шаблона
extra/Буйко_ЛР3_референс.docx: абзацы шаблона копируются, текст заменяется.
Разделы требований — по шаблону улучшенного набора требований из книги С. Куликова.
Содержание собирается в два прохода: номера страниц берутся из PDF первого прохода.
"""

import copy
import os
import re
import shutil
import subprocess

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

WORKSPACE_DIR = "/home/pavel/My/3kurs/RiATPo/Lab3"
UNIVER_DIR = "/home/pavel/Univer/РиАТПО/ЛР3_Доработка_требований"
TEMPLATE_PATH = os.path.join(WORKSPACE_DIR, "extra", "Буйко_ЛР3_референс.docx")
MD_REPORT_PATH = os.path.join(WORKSPACE_DIR, "README.md")
DOCX_REPORT_PATH = os.path.join(WORKSPACE_DIR, "Забелич_451001_3.docx")
PDF_REPORT_PATH = os.path.join(WORKSPACE_DIR, "Забелич_451001_3.pdf")

# Индексы абзацев шаблона
T_FACULTY, T_DISCIPLINE = 7, 9
T_SIGN_1, T_SIGN_2, T_SIGN_3 = 26, 27, 28
T_TITLE_LAST = 35      # «Минск 2026»
T_PAGE_BREAK = 36
T_DOC_TITLE = 37       # «Требования к разрабатываемому приложению»: по центру, полужирный
T_SECTION = 38         # пункт 1., 2., ... верхнего уровня, полужирный
T_SECTION_TEXT = 39    # обычный абзац с отступом первой строки
T_SUBSECTION = 74      # «8.1. ...»: полужирный, без нумерации списка
T_BULLET = 66          # маркер «o»
T_EMPTY = 123
# Первые пункты списков a), b), ... — у каждого списка своя нумерация, чтобы буквы начинались с a)
T_LETTER_LISTS = [41, 44, 47, 52, 59, 65, 75, 78, 81, 84, 87, 91, 95, 103, 109, 120]

# ============================================================================
# СОДЕРЖАНИЕ ОТЧЁТА
# ("sec", номер, заголовок) | ("sub", "8.1.", заголовок) | ("p", текст)
# ("list", [пункт | (пункт, [подпункты-маркеры])])
# ============================================================================

CONTENT = [
    ("sec", 1, "Суть проекта"),
    ("p", "Приложение «File Searcher» (далее — приложение) — настольная утилита для ОС Windows, "
          "предназначенная для поиска аудио-, видео- и офисных файлов в локальных и сетевых каталогах "
          "по категории или маске и вывода сведений о найденных файлах в таблицу."),

    ("sec", 2, "Основные цели проекта"),
    ("list", [
        "Сократить время, которое пользователь тратит на поиск медиафайлов и документов на дисках "
        "компьютера и в сетевых папках.",
        "Обеспечить работу на рабочих станциях под Windows XP и Windows 7 без установки "
        "дополнительного программного обеспечения.",
    ]),

    ("sec", 3, "Критерии достижения целей проекта"),
    ("list", [
        "Поиск по каталогу, содержащему 50 000 файлов, на эталонном стенде (п. 8.1) занимает "
        "не более 100 секунд.",
        "Приложение запускается на чистых ОС Windows XP SP3 и Windows 7 SP1 без установки "
        "библиотек и без прав администратора.",
    ]),

    ("sec", 4, "Риски"),
    ("list", [
        "Искажение кириллических имён файлов: строки Delphi 7 по умолчанию однобайтовые (ANSI).",
        "Зависание поиска при недоступной сетевой папке или обрыве сети.",
        "Зацикливание обхода на точках соединения (Junction) и символических ссылках NTFS.",
    ]),

    ("sec", 5, "Системные характеристики"),
    ("list", [
        "Операционная система: Windows XP SP3, Windows 7 SP1 (32- и 64-разрядные).",
        ("Аппаратные требования:", [
            "Процессор: x86 с частотой от 1 ГГц.",
            "Оперативная память: не менее 256 МБ свободной ОЗУ.",
            "Дисковое пространство: 10 МБ.",
        ]),
        "Поддерживаемые файловые системы: NTFS, FAT32, exFAT (в Windows XP — с обновлением KB955704) "
        "и сетевые папки по протоколу SMB.",
        "Приложение поставляется одним исполняемым файлом FileSearcher.exe и не требует установки.",
    ]),

    ("sec", 6, "Пользовательские требования"),
    ("list", [
        "Пользователь указывает каталог или список каталогов для поиска вводом пути в поле "
        "«Где искать» либо выбором папки кнопкой «…».",
        "Пользователь выбирает категорию файлов в комбо-боксе «Что искать» или вводит собственную маску.",
        "Пользователь запускает поиск кнопкой «Искать» и досрочно прекращает его кнопкой «Остановить».",
        "Пользователь просматривает найденные файлы в таблице, сортирует её по столбцам и видит первый "
        "кадр выбранного видеофайла в панели «Предпросмотр».",
        "Пользователь переключает язык интерфейса (русский или английский) через меню «Language» "
        "без перезапуска приложения.",
    ]),

    ("sec", 7, "Бизнес-правила"),
    ("list", [
        ("Категории включают следующие расширения:", [
            "Аудиофайлы: .mp3, .ogg, .wav, .mid.",
            "Видеофайлы: .avi, .mpg, .mpeg.",
            "Офисные файлы: .doc, .docx, .xls, .xlsx.",
        ]),
        "Расширения и маски сравниваются без учёта регистра букв.",
        "Глубина обхода — не более 64 уровней вложенности от начального каталога; точки соединения "
        "и символические ссылки не раскрываются.",
        "Каталоги и файлы без прав на чтение, тома с неподдерживаемой файловой системой и пути длиннее "
        "260 символов пропускаются с записью WARNING в журнал, поиск продолжается.",
        "Сетевая папка, не ответившая в течение 10 секунд, пропускается с записью в журнал.",
        "Лимит времени поиска по умолчанию — 60 минут. По его истечении поиск приостанавливается, "
        "и приложение предлагает пользователю продолжить или завершить поиск.",
    ]),

    ("sec", 8, "Атрибуты качества"),
    ("sub", "8.1.", "Производительность"),
    ("list", [
        "На эталонном стенде (Intel Core i3, 4 ГБ ОЗУ, HDD SATA 7200 об/мин, NTFS) приложение "
        "проверяет не менее 500 файлов в секунду на локальном диске.",
        "Поиск выполняется в отдельном потоке; время отклика окна на действия пользователя во время "
        "поиска не превышает 0,5 секунды.",
    ]),
    ("sub", "8.2.", "Надёжность"),
    ("list", [
        "При ошибке доступа, повреждённом файле или обрыве сети приложение пропускает объект, "
        "записывает событие в журнал и продолжает поиск без аварийного завершения.",
        "После остановки поиска (вручную или по лимиту времени) найденные файлы остаются в таблице.",
    ]),
    ("sub", "8.3.", "Удобство использования"),
    ("list", [
        "Все элементы управления доступны с клавиатуры, поиск запускается клавишей Enter из поля "
        "«Где искать».",
        "Текущий проверяемый каталог отображается в строке состояния внизу окна в режиме реального времени.",
    ]),
    ("sub", "8.4.", "Локализация"),
    ("list", [
        "Все надписи и сообщения интерфейса переведены на русский и английский языки.",
        "При первом запуске язык выбирается по языку ОС; если язык ОС не русский, используется английский.",
        "Новый язык добавляется файлом перевода в папке Lang без перекомпиляции (п. 10.5).",
    ]),

    ("sec", 9, "Ограничения"),
    ("list", [
        "Приложение разрабатывается на Delphi 7; для работы с именами файлов используются "
        "Unicode-функции Win32 API (FindFirstFileW, FindNextFileW).",
        "Поиск выполняется только по именам файлов, поиск по содержимому не поддерживается.",
        "Приложение не открывает и не воспроизводит найденные файлы.",
        "Разделы UNIX (ext2/3/4) доступны только как сетевые папки; FTP, WebDAV и облачные "
        "хранилища не поддерживаются.",
    ]),

    ("sec", 10, "Детальные спецификации"),
    ("sub", "10.1.", "Выбор каталогов"),
    ("list", [
        "Поле «Где искать» принимает до 1024 символов; пути разделяются символом «;», пробелы по краям "
        "путей удаляются.",
        "Допускаются UNC-пути сетевых папок (\\\\server\\share) и локальные пути (C:\\Music).",
        "Кнопка «…» открывает стандартный диалог выбора папки Windows и добавляет выбранный путь "
        "в конец строки поля.",
        "Перед поиском приложение проверяет существование каждого пути: несуществующие пути пропускаются "
        "с сообщением, а если корректных путей нет, поиск не запускается.",
    ]),
    ("sub", "10.2.", "Параметры поиска"),
    ("list", [
        ("Комбо-бокс «Что искать» содержит пункты:", ["«Аудиофайлы»;", "«Видеофайлы»;", "«Офисные файлы»."]),
        "Поле комбо-бокса редактируемое: пользователь вводит маску со знаками * и ?, маски разделяются "
        "символом «;» (например, *.pdf;report*.*).",
        "Лимит времени поиска задаётся в минутах в файле настроек FileSearcher.ini; значение 0 отключает лимит.",
    ]),
    ("sub", "10.3.", "Процесс поиска"),
    ("list", [
        "Нажатие кнопки «Искать» запускает поиск; на время поиска надпись кнопки меняется на "
        "«Остановить», поля ввода блокируются.",
        "Нажатие кнопки «Остановить» прекращает поиск не позднее чем через 1 секунду.",
        "По истечении лимита времени выводится диалог «Поиск длится больше N мин. Продолжить?» "
        "с кнопками «Продолжить» и «Завершить».",
        "Счётчик «Найдено файлов» обновляется по мере обнаружения файлов.",
    ]),
    ("sub", "10.4.", "Отображение результатов"),
    ("list", [
        ("Таблица результатов содержит столбцы:", [
            "№ (порядковый номер).",
            "Имя (имя файла с расширением).",
            "Путь (полный путь к каталогу).",
            "Размер (в Б, КБ, МБ или ГБ с одним знаком после запятой).",
            "Создан (дата-время создания в формате ДД.ММ.ГГГГ ЧЧ:ММ).",
            "Изменён (дата-время последнего изменения в том же формате).",
        ]),
        "Щелчок по заголовку столбца сортирует таблицу по возрастанию, повторный щелчок — по убыванию.",
        "При выборе строки с видеофайлом панель «Предпросмотр» справа от таблицы показывает его первый "
        "кадр; для аудио- и офисных файлов, а также при ошибке декодирования показывается значок типа файла.",
        "Строка состояния внизу окна выводит текст «Сейчас проверяется: <путь>».",
    ]),
    ("sub", "10.5.", "Журналирование и файлы перевода"),
    ("list", [
        "Приложение ведёт журнал FileSearcher.log в каталоге программы; каждая запись содержит "
        "дату-время, уровень (INFO, WARNING, ERROR) и текст события.",
        "При достижении размера 1 МБ (1 048 576 байт) запись в журнал прекращается, последней добавляется "
        "запись «Достигнут предельный размер журнала», в строке состояния выводится предупреждение.",
        "При следующем запуске прежний журнал переименовывается в FileSearcher.old.log, и запись "
        "начинается в новый файл.",
        "Файл перевода — текстовый файл UTF-8 со строками вида ключ=значение, имя файла — код языка "
        "(ru.lng, en.lng).",
    ]),
]

DOC_TITLE = "Требования к разрабатываемому приложению"


def toc_entries():
    """(ключ, текст в содержании, уровень, строка для поиска в PDF)."""
    out = [("title", DOC_TITLE, 0, DOC_TITLE)]
    for block in CONTENT:
        if block[0] == "sec":
            text = f"{block[1]}. {block[2]}"
            out.append((text, text, 0, block[2]))
        elif block[0] == "sub":
            text = f"{block[1]} {block[2]}"
            out.append((text, text, 1, text))
    return out


# ============================================================================
# Работа с XML шаблона
# ============================================================================

def strip_ids(el):
    """Убирает w14:paraId/textId, чтобы скопированные абзацы не дублировали идентификаторы."""
    for node in el.iter():
        for attr in list(node.attrib):
            if attr.endswith("}paraId") or attr.endswith("}textId"):
                del node.attrib[attr]
    return el


def first_run_rpr(p_el):
    r = p_el.find(qn("w:r"))
    if r is not None and r.find(qn("w:rPr")) is not None:
        return copy.deepcopy(r.find(qn("w:rPr")))
    return parse_xml(f'<w:rPr {nsdecls("w")}><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" '
                     f'w:cs="Times New Roman"/><w:sz w:val="28"/><w:szCs w:val="28"/></w:rPr>')


def make_run(rpr, text, bold=False):
    r = parse_xml(f'<w:r {nsdecls("w")}/>')
    rpr = copy.deepcopy(rpr)
    for tag in ("w:b", "w:bCs", "w:i", "w:iCs", "w:rStyle", "w:highlight", "w:shd"):
        for e in rpr.findall(qn(tag)):
            rpr.remove(e)
    if bold:
        rpr.insert(1, parse_xml(f'<w:bCs {nsdecls("w")}/>'))
        rpr.insert(1, parse_xml(f'<w:b {nsdecls("w")}/>'))
    r.append(rpr)
    for i, part in enumerate(text.split("\t")):
        if i:
            r.append(parse_xml(f'<w:tab {nsdecls("w")}/>'))
        t = parse_xml(f'<w:t {nsdecls("w")} xml:space="preserve"></w:t>')
        t.text = part
        r.append(t)
    return r


def set_runs(p_el, segments, rpr=None):
    """Заменяет содержимое абзаца прогонами [(текст, полужирный), ...] с оформлением первого прогона."""
    rpr = rpr if rpr is not None else first_run_rpr(p_el)
    for child in list(p_el):
        if child.tag != qn("w:pPr"):
            p_el.remove(child)
    for text, bold in segments:
        if text:
            p_el.append(make_run(rpr, text, bold))
    return p_el


def clone(p_el, segments):
    return set_runs(strip_ids(copy.deepcopy(p_el)), segments)


def keep_with_next(p_el):
    """Не отрывать заголовок от следующего абзаца."""
    pPr = p_el.find(qn("w:pPr"))
    if pPr.find(qn("w:keepNext")) is None:
        pPr.insert(1 if pPr.find(qn("w:pStyle")) is not None else 0, parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
    return p_el


def add_page_numbers(doc):
    """Номер страницы в нижнем колонтитуле справа, на титульном листе не выводится."""
    sect = doc.sections[0]
    sect.different_first_page_header_footer = True
    p = sect.footer.paragraphs[0]
    p.alignment = 2
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    rpr = ('<w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'
           '<w:sz w:val="28"/></w:rPr>')
    for xml in (f'<w:r {nsdecls("w")}>{rpr}<w:fldChar w:fldCharType="begin"/></w:r>',
                f'<w:r {nsdecls("w")}>{rpr}<w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>',
                f'<w:r {nsdecls("w")}>{rpr}<w:fldChar w:fldCharType="separate"/></w:r>',
                f'<w:r {nsdecls("w")}>{rpr}<w:t>2</w:t></w:r>',
                f'<w:r {nsdecls("w")}>{rpr}<w:fldChar w:fldCharType="end"/></w:r>'):
        p._p.append(parse_xml(xml))


# ============================================================================
# Сборка документа
# ============================================================================

def build_document(toc_pages=None):
    toc_pages = toc_pages or {}
    doc = docx.Document(TEMPLATE_PATH)
    P = [p._p for p in doc.paragraphs]

    tpl = {k: copy.deepcopy(P[i]) for k, i in {
        "page_break": T_PAGE_BREAK, "doc_title": T_DOC_TITLE, "section": T_SECTION,
        "section_text": T_SECTION_TEXT, "subsection": T_SUBSECTION, "bullet": T_BULLET,
        "empty": T_EMPTY}.items()}
    letter_lists = [copy.deepcopy(P[i]) for i in T_LETTER_LISTS]

    # --- Титульный лист: правим текст абзацев шаблона ---
    set_runs(P[T_FACULTY], [("Факультет компьютерных систем и сетей", False)])
    set_runs(P[T_DISCIPLINE], [("Дисциплина: Разработка и анализ требований к программному обеспечению", False)])
    sign_1 = doc.paragraphs[T_SIGN_1].text
    gap = sign_1.index("Проверил:") - len("ст. гр. 451001")
    set_runs(P[T_SIGN_2], [("ст. гр. 451001" + " " * gap + "Данилова Г. В.", False)])
    set_runs(P[T_SIGN_3], [("Забелич П. А.", False)])

    # --- Удаляем всё содержимое шаблона после титульного листа ---
    body = doc.element.body
    started = False
    for child in list(body):
        if started and child.tag != qn("w:sectPr"):
            body.remove(child)
        if child is P[T_TITLE_LAST]:
            started = True
    sect_pr = body.find(qn("w:sectPr"))

    def add(el):
        sect_pr.addprevious(el)

    def page_break():
        add(strip_ids(copy.deepcopy(tpl["page_break"])))

    # --- Содержание ---
    page_break()
    add(clone(tpl["doc_title"], [("СОДЕРЖАНИЕ", True)]))
    for key, text, level, _ in toc_entries():
        p = clone(tpl["empty"], [(f"{text}\t{toc_pages.get(key, 3)}", False)])
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            pPr = parse_xml(f'<w:pPr {nsdecls("w")}/>')
            p.insert(0, pPr)
        pPr.append(parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="9355"/></w:tabs>'))
        pPr.append(parse_xml(f'<w:spacing {nsdecls("w")} w:after="40" w:line="276" w:lineRule="auto"/>'))
        pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="{0 if level == 0 else 567}" w:right="0"/>'))
        add(p)

    # --- Требования ---
    page_break()
    add(clone(tpl["doc_title"], [(DOC_TITLE, True)]))
    lists = iter(letter_lists)
    for block in CONTENT:
        kind = block[0]
        if kind == "sec":
            add(keep_with_next(clone(tpl["section"], [(block[2], True)])))
        elif kind == "sub":
            add(keep_with_next(clone(tpl["subsection"], [(f"{block[1]} {block[2]}", True)])))
        elif kind == "p":
            add(clone(tpl["section_text"], [(block[1], False)]))
        elif kind == "list":
            list_tpl = next(lists)
            for item in block[1]:
                text, bullets = (item, []) if isinstance(item, str) else item
                add(clone(list_tpl, [(text, False)]))
                for b in bullets:
                    add(clone(tpl["bullet"], [(b, False)]))

    add_page_numbers(doc)

    cp = doc.core_properties
    cp.author = "Забелич П. А."
    cp.last_modified_by = "Забелич П. А."
    cp.title = "Лабораторная работа №3. Доработка требований. Вариант 3"
    cp.comments = ""
    return doc


# ============================================================================
# Конвейер: два прохода, PDF, README
# ============================================================================

def convert_to_pdf():
    soffice_bin = "/opt/libreoffice26.2/program/soffice"
    if not os.path.exists(soffice_bin):
        soffice_bin = "soffice"
    res = subprocess.run([soffice_bin, "--headless", "--convert-to", "pdf", DOCX_REPORT_PATH,
                          "--outdir", WORKSPACE_DIR], capture_output=True, text=True)
    if res.returncode != 0:
        print("LibreOffice PDF conversion stderr:", res.stderr)


def find_pages_in_pdf(pdf_path):
    res = subprocess.run(["pdftotext", pdf_path, "-"], capture_output=True, text=True)
    pages = [re.sub(r"\s+", " ", p) for p in res.stdout.split("\x0c")]
    mapping = {}
    for key, _, _, needle in toc_entries():
        for p_idx in range(2, len(pages)):  # пропускаем титульный лист и содержание
            if needle in pages[p_idx]:
                mapping[key] = p_idx + 1
                break
        else:
            print(f"WARNING: heading not found in PDF: {needle}")
    print("Detected page mapping:", mapping)
    return mapping


def generate_markdown():
    out = [
        "Министерство образования Республики Беларусь  ",
        "Учреждение образования «Белорусский государственный университет информатики и радиоэлектроники»  ",
        "Факультет компьютерных систем и сетей  ",
        "Кафедра программного обеспечения информационных технологий  ",
        "Дисциплина: Разработка и анализ требований к программному обеспечению",
        "",
        "**ОТЧЕТ**  ",
        "Лабораторная работа №3  ",
        "Вариант 3",
        "",
        "Выполнил: ст. гр. 451001 Забелич П. А.  ",
        "Проверил: Данилова Г. В.",
        "",
        "Минск 2026",
        "",
        "---",
        "",
        f"# {DOC_TITLE}",
    ]
    for block in CONTENT:
        kind = block[0]
        if kind == "sec":
            out += ["", f"## {block[1]}. {block[2]}", ""]
        elif kind == "sub":
            out += ["", f"### {block[1]} {block[2]}", ""]
        elif kind == "p":
            out.append(block[1])
        elif kind == "list":
            for letter, item in zip("abcdefghij", block[1]):
                text, bullets = (item, []) if isinstance(item, str) else item
                out.append(f"{letter}) {text}")
                out += [f"   - {b}" for b in bullets]
    with open(MD_REPORT_PATH, "w", encoding="utf-8") as f:
        f.write(re.sub(r"\n{3,}", "\n\n", "\n".join(out)) + "\n")
    print(f"Generated Markdown: {MD_REPORT_PATH}")


def sync_to_univer():
    if os.path.exists(UNIVER_DIR):
        for f in [DOCX_REPORT_PATH, PDF_REPORT_PATH, MD_REPORT_PATH]:
            shutil.copy(f, UNIVER_DIR)
        print(f"Synchronized report files to {UNIVER_DIR}")


def run_pipeline():
    print("--- PASS 1: draft ---")
    build_document().save(DOCX_REPORT_PATH)
    convert_to_pdf()
    pages_map = find_pages_in_pdf(PDF_REPORT_PATH)

    print("--- PASS 2: exact TOC pages ---")
    build_document(pages_map).save(DOCX_REPORT_PATH)
    convert_to_pdf()
    generate_markdown()
    sync_to_univer()
    print("All tasks completed successfully!")


if __name__ == "__main__":
    run_pipeline()
